from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path

ROOT=Path(__file__).resolve().parent.parent
doc=Document()
s=doc.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(2); s.bottom_margin=Cm(1.8); s.left_margin=Cm(2); s.right_margin=Cm(2)
s.header_distance=Cm(.8); s.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Microsoft YaHei'; st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    st.font.size=Pt(10.5)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
header=s.header.paragraphs[0]; header.text='项目研究  /  展会排期与需求变更'; header.style='Normal'
header.runs[0].font.size=Pt(8)
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('调研截止 2026年10月23日    |    ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)

def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,2)
def page(t):
    par=doc.add_heading(t,1)
    par.paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit'); trPr.append(no)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'193B52' if ri==0 else ('F0F4F7' if ri%2 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'85');m.set(qn('w:type'),'dxa');margins.append(m)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.12
                for r in par.runs:
                    r.font.size=Pt(9)
                    if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    gap=p('');gap.paragraph_format.space_after=Pt(2);gap.paragraph_format.space_before=Pt(0)
    gap.paragraph_format.line_spacing=Pt(2);gap.paragraph_format.keep_with_next=True
    gap.paragraph_format.widow_control=False
    return t
def link(label,url):
    par=doc.add_paragraph();par.paragraph_format.space_after=Pt(9)
    x=OxmlElement('w:hyperlink');x.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'245B78');pr.append(col);r.append(pr)
    tx=OxmlElement('w:t');tx.text=label;r.append(tx);x.append(r);par._p.append(x)

p('行业现状与开源项目调研', 'Subtitle')
p('展会排期与需求变更\n项目研究报告','Title')
p('截至 2026年10月23日  |  立项与四周试点参考  |  V1.0')
h('核心结论')
p('建议先建设单批次、只读的排期核对与变更追踪工作台：统一申请编号和版本，依据确定性规则识别场地及共享资源的时间冲突，生成可定位来源的缺项与交接清单，正式结论保留人工确认。现有公开证据支持这些能力的技术可行性，但尚不足以证明该企业的实际收益。')
p('本次筛选了5个值得参考的GitHub项目。LibreBooking最接近资源预约与冲突识别；pretalx最适合研究排期版本与发布；OpenProject适合借鉴变更任务与责任追踪；Eventyay覆盖活动运营；FullCalendar提供日历展示。没有核实到一个可直接替代本项目全部流程的现成方案。')
h('证据边界')
p('本地输入 discovery-report.md 明确声明：人物、指标和业务问题来自合成场景。文中“每月45单、每单35分钟”、36/36调查覆盖、四周试点和30样本均应按原文性质解读；前两项不能当作真实经营基线，覆盖率不能当作验收结果。[L1]')
p('行业部分采用中国贸促会、上海官方公开资料和厂商一手文档；项目部分核对仓库元数据、README、许可证及两个核心源码文件。属于桌面研究，未部署软件、未执行端到端测试、未访谈真实客户，也未取得产品报价。')
h('阅读指引')
table(['章节','内容'],[['一至二','行业规模、区域动态、商业产品能力与项目机会'],['三至五','GitHub对比、源码证据、选型与复用建议'],['六至八','需求映射、数据架构、冲突和变更规则'],['九至十','四周试点、验收指标、待确认事项'],['参考资料','可点击的原始来源与核验说明']],[3,14])

page('一 行业现状与项目背景')
p('可核实的全国完整年度数据以2025年为主，2026年补充已发布的地方政策和展会案例。本报告不将2026年局部活动外推为全年规模，也不把展览面积换算为软件市场收入。')
table(['观察项','公开证据','对项目的启示'],[
['国内展览规模','2025年经贸类展会4095项，展览总面积1.59亿平方米，分别同比增长6.53%和2.5%。[S1]','排期协作具有明确业务场景；不能据此推算本企业单量。'],
['区域集聚','京津冀、长三角、珠三角合计贡献全国53.85%的展会数量和62.48%的展出面积。[S1]','上海项目适合优先观察多场地及跨团队交接。'],
['上海产业方向','2026至2028年行动方案提出建立MICE活动日历与信息共享机制，提前发布重点活动年度安排。[S2]','应区分对外活动日历与内部资源占用及审批台账。'],
['2026年活动样本','上海官方披露第十二届上交会于6月11至13日举办，面积3.8万平方米、参展企业1006家。[S3]','大型展会同时涉及展区及配套活动；此案例不证明本企业发生过冲突。']
],[2.5,7.2,7.3])
h('从行业趋势到本项目机会')
p('中国贸促会将数字化与绿色化列为行业发展特征。[S1] 对本项目而言，最直接的数字化切入点是把时间窗、区域、共享资源、版本和确认状态形成一致记录。是否需要数字孪生、智能推荐或跨场馆优化，应由实际复杂度和数据完整性决定。')
p('上海行动方案的活动日历有利于公共信息协同，但城市级日历无法替代内部装卸口、搭建、撤场、恢复检查的分钟级资源核对。这是基于公开政策与本地场景的分析推断，不是政策对软件功能的强制要求。[S2][L1]')
h('研究限制')
p('本轮未取得可直接核验的UFI最新报告全文，因此不引用其2026年预测比例。也未找到能支持“中国会展企业普遍依赖群消息和Excel”这一普及率判断的一手样本；相关问题仅保留为本地报告中的待验证假设。')

page('二 商业产品能力与竞争位置')
p('商业产品已覆盖排期、预约审批和运营协作。项目价值应落在本地交接规则、来源追溯和变更影响解释，而非仅增加一个日历页面。以下是厂商公开能力描述，不代表独立实施效果评测。')
table(['产品或路径','已核实的公开能力','与本项目的关系'],[
['Momentus Enterprise','支持多日多空间活动、占档与释放、变更追踪及审批、跨部门角色权限。[S4]','业务覆盖最接近完整场馆运营平台；需核查具体版本、导出接口与配置成本。'],
['Skedda或AllBooked','按空间、用户标签或日期设置审批；请求可并存且不占用空间，审批后形成预约。[S5]','可借鉴“请求不等于确认”；不能直接替代本项目占档状态口径。'],
['开源资源预约','LibreBooking支持多资源预约、候补、角色权限、使用报告。[G1]','可用于验证资源建模和时间冲突；变更确认链仍需适配。'],
['定制只读工作台','按现有台账字段做版本比对、来源定位和提示。属于本报告建议。','最贴合当前只读边界，需单独承担规则、审计与维护。']
],[3,7,7])
h('商业方案应如何核验')
p('Momentus把合同状态、空间状态、服务订单等纳入业务平台；从产品方向看可覆盖更长链条。[S4] 但本轮无权访问实际租户，无法判断接口可用性、国产化适配、部署地域、迁移工时和报价。采购前应使用本项目30样本做现场演示，要求逐条展示来源、版本、冲突原因及人工确认记录。')
p('Skedda的待审批请求不进入已确认排期，多个冲突请求可以同时存在；审批功能也受套餐限制。[S5] 本项目应在内部工作台同时看见“意向、待确认、已确认”，并明确冲突等级，避免把看不见请求误解成没有冲突。')
h('建议的差异化定位')
p('把已有服务台、业务台账与文件库的只读导出连接起来，为同一申请建立一条可查证的变更链。优先解决“拿到哪个版本、哪些字段没确认、影响哪些资源、下一位谁接收”，再评估是否替换现有运营平台。[L1]')

page('三 GitHub候选项目总览')
p('检索覆盖 venue booking、exhibition management、pretalx，并按资源排期、版本发布、变更追踪和展示组件补充候选。通用运动场预约、纯界面示例、景区展示系统等搜索结果未进入主推荐。筛选是目的性抽样，不是GitHub全部项目排名。')
table(['项目','用途与主语言','许可证核查','最近推送日期'],[
['LibreBooking/librebooking','资源预约\nPHP','GPL-3.0','2026-10-22'],
['pretalx/pretalx','会议议程与发布\nPython','当前LICENSE为AGPLv3并附加条款；README仍写Apache','2026-10-23'],
['opf/openproject','项目及工作项管理\nRuby','GPL-3.0','2026-10-22'],
['fossasia/eventyay','活动运营平台\nPython','仓库元数据标识Apache-2.0；依赖另核','2026-10-22'],
['fullcalendar/fullcalendar','日历组件\nTypeScript','Standard为MIT；Premium另行授权','2026-10-16']
],[4.8,3.5,5.5,3.2])
p('元数据为2026年10月23日查询快照，日期按GitHub返回的UTC日期记录。五个仓库查询时均未归档。“最近推送”可能包含自动更新或开发分支改动，不等于稳定版发布时间，也不证明生产可用性。[G1至G5]')
h('功能适配判断')
table(['项目','最值得借鉴','本项目必须补齐'],[
['LibreBooking','资源占用、缓冲、预约冲突','有效版本、来源行、跨团队变更链'],
['pretalx','草稿、版本发布、历史差异','场馆分区及装卸口资源规则'],
['OpenProject','责任人、状态与工作项协作','场地时间冲突与只读数据导入'],
['Eventyay','活动、议程及插件组织','场馆恢复检查和合同前交接口径'],
['FullCalendar','月周日视图、前端集成','全部后台业务规则及审计']
],[3,6.2,7.8])
p('“必须补齐”指本轮证据未证明其原生满足本项目要求，不能直接理解为该项目绝对没有该功能。完整能力仍应通过配置与部署测试确认。')

page('四 资源冲突与版本发布的源码证据')
h('LibreBooking 资源排期的优先参考')
p('README列有多资源预约、候补、角色访问控制与使用报告，部署采用PHP和MySQL或MariaDB，可使用Docker。[G1] 旧仓库地址 LibreBooking/app 已迁移至 LibreBooking/librebooking，选型记录应使用当前地址。')
p('本轮读取 ReservationConflictIdentifier.php，确认其冲突识别遍历预约实例与关联资源，使用资源可用性策略查询时间区间，并考虑缓冲及最大并发数量。代码中存在管理员相关的缓冲处理分支，因此不可默认直接复用其默认行为来满足本项目“人工也不能绕过恢复检查”的规则。[G6]')
p('可迁移的设计思路是“资源集合＋占用时间窗＋并发约束＋可解释冲突结果”；需要额外实现版本有效性、原始行定位、装卸口与展馆区域联动、待确认值区分。首轮建议研究模型与规则，不将完整预约写入链路接到业务系统。')
p('当前README明确提示7.1.0以前版本存在严重登录安全问题，并建议升级。若做部署验证，应选择修复后的正式版本并复核其依赖；本报告未安装或扫描，因此不宣称安全验收通过。[G1]')
h('pretalx 版本机制的优先参考')
p('README说明支持多房间、多日议程的构建、版本化与发布。schedule.py 定义一个活动拥有一份未发布排期及多份已发布排期，并保存版本号、发布时间、说明，以及前一版本与差异入口。[G2][G7] 这些结构很适合借鉴到“已确认基线＋变更草稿＋确认后形成新基线”。')
p('适用限制是领域不同：演讲议程排期不能直接等同于场馆占档、合同审核或进撤场计划。应保留展会、场馆区域、共享资源与业务审批记录之间的独立关系。')
h('授权核验中的实际分歧')
p('查询当日，pretalx README末尾写有Apache授权，但根目录LICENSE及本轮读取的schedule.py文件头指向AGPLv3及附加条款。[G2][G7][G8] 本报告按具体LICENSE和文件声明记录，不将其归类为可无条件复制的Apache项目。若拟复制源码，先锁定版本并复核对应文件授权；当前建议优先借鉴设计。')

page('五 变更协作与展示组件选型')
h('OpenProject 适合参考变更工单')
p('官方README列有项目计划、甘特图、任务、缺陷、协作及自托管社区版，也区分企业扩展。[G3] 可把一次需求变更映射为工作项，关联申请、接收岗位、截止时间和处置结果。但甘特图的任务依赖不能替代区域占用的时间冲突计算，完整访问审计与审批配置尚需实测。')
h('Eventyay 适合参考活动平台组织方式')
p('当前仓库已统一票务、征集、讲者、议程、视频等能力，README列有REST接口、插件和展览相关扩展。[G4] 本轮读取的开发分支采用Python、Django、PostgreSQL、Redis及Celery，覆盖面较大。若只做只读核对，直接部署整个平台会带来不必要的配置与维护工作；本轮未验证展览插件对场馆排期的覆盖深度。')
h('FullCalendar 适合用作展示层')
p('仓库提供可拖动日历及React、Vue、Angular等集成入口。[G5] 项目首版宜关闭直接拖动保存，点击日程时打开申请版本与来源。资源时间轴等功能应逐插件确认是否属于Premium，不能把Standard的MIT许可证扩展到所有高级组件。[S6]')
table(['实施路径','本轮适配判断','建议'],[
['直接部署完整平台','能力广，但字段与流程差异明显；会超出只读试点','暂不作为首版前提'],
['二次开发预约系统','资源能力较近，仍需适配变更和审计','完成真实样本映射后再评估'],
['定制轻量工作台','最贴合只读、单批次与来源追溯边界','优先；借鉴开源设计'],
['采购商业平台','可覆盖更长运营链路，价格和接口待核','保留对照演示与中长期选项']
],[4,9,4])
p('以上为需求适配判断，不是已完成的性能、成本或安全评分。四周安排只覆盖边界验证及最小原型，不承诺完成企业级系统上线。')

page('六 将调查假设转成产品范围')
p('本节沿用本地报告REQ-1至REQ-6的编号。优先级是试点建议，真实责任人应按岗位确认，不沿用合成人物姓名作为实施分工。[L1]')
table(['需求','首版可交付能力','证据及验证方式'],[
['REQ-1\n版本交接','统一申请编号；展示有效基线、草稿、未结项、接收岗位、截止时间','f01 f20 f24 f23；逐交接点回放'],
['REQ-2\n冲突提示','列出冲突资源、区间、规则、来源与复核状态；不改变占档','f05 f14 f41 f44；与标注结果核对'],
['REQ-3\n只读数据关联','导入脱敏导出文件；保留原行；缺失与未确认分开','f30至f34；检查重复、空值及断链'],
['REQ-4\n跨团队口径','同一申请展示客户使用和恢复检查时间；交接双方确认','f50 f51 f53；同一原始记录对照'],
['REQ-5\n审计留痕','保留查询版本、提示、复核人和结果；读取失败显示快照时间','f35 f43 f45；抽查完整事件链'],
['REQ-6\n样本验收','固定30个标注样本、规则版本与回放结果；记录误报漏报','f02 f03 f42；业务与质量共同复核']
],[3,8,6])
h('首版边界')
p('输入为一个批次的脱敏档期表、变更单及经授权的资源分区资料；输出为排期视图、冲突清单、版本差异和一页交接摘要。应用可以在独立试点库保存导入快照及复核记录，但不写回原业务系统，不调用外部消息、签约、锁档或执行接口。')
p('排除项包括自动承诺容量、自动优选商业档期、合同签署、资金结算、真实客户通知和跨企业资料自由访问。资源容量及可共享规则不完整时转为“待核查”，不能因无法计算而输出“无冲突”。')
h('四来源收敛结果')
p('Buyer与Operator目前仍是合成陈述，Brief缺少真实书面版本；本轮新增Industry证据，证明市场已有资源预约、版本发布和协作工具。这只补足了行业可行性视角，仍不能把整个需求标为“已被真实客户验证”。')

page('七 建议架构与核心数据')
p('推荐链路：只读导出 → 原始快照留存 → 字段标准化与质量检查 → 版本筛选 → 规则计算 → 工作台展示 → 人工复核记录 → 只读报告导出。首轮可以批量导入，不依赖生产接口。')
table(['对象','最小字段','关键约束'],[
['申请 Application','application_id、source_id、event_id、batch_id','内部ID与来源编号分开；缺少统一编号时建立经人工确认的映射'],
['版本 Version','version_id、parent_id、状态、有效区间、发生时间、接收时间','原始快照不覆盖；乱序记录不按接收时间自动晋级'],
['资源占用 Occupancy','resource_id、资源类型、开始结束、用途、容量口径','区域、装卸口、通道分别建模；同一展会可关联多条占用'],
['变更 Change','change_id、基线版本、字段旧新值、原因、受影响对象','草稿不能替换已确认基线；新基线确认时校验父版本'],
['交接 Handover','发送岗位、接收岗位、未结项、截止时间、确认记录','未确认责任不自动转移；保留退回及撤回原因'],
['证据与审计 Evidence','来源文件标识、原行、快照哈希、规则版本、查询与复核事件','结果能够回到输入版本；保存访问与决策过程']
],[3,7.1,6.9])
h('有效版本如何确定')
p('先按申请映射关联版本，再按业务认可的有效期与确认状态选择基线；收到新文件不代表已经生效。对同一时间存在多个已确认有效版本、父版本断链或来源相互矛盾的记录，进入人工队列，不通过文件名或最大时间戳猜测。')
h('技术实现建议')
p('如团队无既定技术约束，可采用常用Web框架、关系数据库和独立规则模块，前端用标准日历组件。30样本试点先实现可回放的区间比较，无需优化求解器。AI可选用于提取变更描述及生成摘要，提取值标记待确认；日期运算、版本选择与正式结论使用确定性逻辑及人工复核。')
p('读取失败时保留上次快照并显示抓取时间和失败状态；原始文件与派生结果按权限访问。留存与删除期限需由真实项目负责人确认后配置，本报告不虚构期限。')

page('八 冲突判定与变更流程')
h('先定义占用时间再计算冲突')
p('每个资源记录自己的占用区间，可分别包含进场搭建、客户使用、撤场、恢复检查。采用左闭右开区间［开始，结束），同一不可共享资源的两个有效区间满足 max(开始A，开始B) < min(结束A，结束B) 时提示重叠。端点相接是否需要缓冲，必须先按经确认的规则扩展区间。')
table(['规则','提示内容','边界处理'],[
['区域互斥','同一展区或不可同时使用的分区发生重叠','需真实分区互斥表；不能仅比名称'],
['共享资源','不同展会使用同一装卸口或通道且重叠','容量未知时提示待核查；不推断承载能力'],
['恢复缓冲','客户延时导致撤场或恢复检查不足','缓冲值由运营确认，不硬编码通用分钟数'],
['版本异常','过期来源、双有效版本或变更基线已变动','归入数据或版本异常，不混算物理冲突'],
['缺项异常','时间、资源、确认状态缺失或单位不一致','进入人工队列，不默认无冲突']
],[3,7,7])
h('变更的处理链')
p('提交变更草稿 → 绑定当前有效基线 → 展示字段差异及受影响资源 → 重算冲突 → 指定岗位复核 → 授权人员确认或退回 → 生成新有效版本与交接清单。试点中的“确认”只记录离线复核结论，不触发原系统占档、合同或对外承诺。')
p('若复核期间另一份变更已形成新基线，本次草稿必须重新比对并复核；不得沿用旧基线的“无冲突”结论。撤回或作废同样保留原因和原记录。')
h('合成示例 用于说明规则')
p('活动A使用装卸口D的撤场区间为10月26日16时至18时；活动B的进场为17时30分至19时，规则应提示17时30分至18时重叠。若A把撤场延长到18时30分，则重新生成差异并计算重叠至18时30分。输出还需包含两份申请及版本、来源行、规则编号和待复核岗位。此例不是实际运营事件。')

page('九 四周试点安排与资源')
p('原报告给定四周、两名业务骨干、每周投入不超过四小时，但未说明四小时为每人还是合计。[L1] 在确认前按业务协作每周合计不超过四小时规划；这不是开发预算。实施负责人、开发投入和基础设施费用需另行落实。')
table(['阶段','工作与可审查产物','业务投入建议'],[
['第1周\n补齐口径','确认一个批次、书面Brief、授权范围、岗位映射；形成字段字典、冲突规则及30样本标注集','合计4小时内：样本与规则评审'],
['第2周\n建立基线','完成导入、原行定位、版本与质量检查；交付首轮数据质量清单和有效版本列表','合计4小时内：核对字段与版本'],
['第3周\n原型回放','完成规则提示、版本差异、交接摘要；保存首轮回放与误报漏报原因','合计4小时内：复核异常及调整规则'],
['第4周\n冻结复验','冻结规则版本，复跑同一批样本；提交验收记录、已知限制及下一阶段建议','合计4小时内：联合验收评审']
],[3,10,4])
h('30样本建议构成')
p('建议8例区域或分区冲突、6例共享资源冲突、4例恢复缓冲冲突、4例无冲突对照、4例时间或字段缺项、4例乱序版本及撤回场景，共30例。每例按主要测试目的归类，可有多个标签；这是建议配比，需以真实业务类型修订，不代表已有样本。')
p('由运营岗位给出预期结果和规则依据，质量岗位复核，数字化岗位保存输入版本和回放记录。发生分歧的样本先补充判定口径，不把系统结果当作标注答案。应保留若干未参与规则调整的样本，避免只验证已经反复调好的案例。')
h('阶段准入与停止条件')
p('未获得脱敏样本授权、找不到正式确认岗位、时间口径无法对齐时，停留在合成数据演练。出现越权访问、旧版本覆盖新版本、来源不可追溯或自动执行业务动作时暂停试点并处理原因。质量岗位保留停止试点权利，真实岗位与升级方式应写入项目责任清单。')

page('十 验收口径与立项待办')
table(['指标','计算与记录口径','本轮建议'],[
['已标注冲突召回','正确提示的标注冲突数 ÷ 全部标注冲突数','30样本内识别全部已标注冲突；不外推为生产100%'],
['误报与漏报','误报按非标注冲突提示逐条复核；漏报登记影响资源与原因','分别披露计数、分母和类别；误报阈值待签字'],
['来源关联完整率','可定位输入版本与原行的结果数 ÷ 全部结果数','建议100%，包括缺项和异常结果'],
['人工复核完整率','有复核人、时间和结论的结果数 ÷ 应复核结果数','建议100%，无确认不能视为闭环'],
['处理工时与返工','同一口径记录实际核对分钟、交接补核次数','先测真实基线，再与辅助处理比较'],
['只读与版本控制','写回及外部动作次数；旧版本误覆盖次数','均为0；至少回放一次读取失败与乱序输入']
],[3.6,8.7,4.7])
p('“全部已标注冲突识别”来自本地演练标准，其余阈值是实施建议。[L1] 冲突数应按资源、规则和版本组合去重，不能用30个样本数直接代替冲突数。无标注冲突的测试集不得报告召回率为100%。')
h('收益评估的正确起点')
p('本地虚构基线45单×35分钟＝1575分钟，即26.25小时／月，只能说明计算方式，不能写入企业收益承诺。真实收益应以授权样本测得的“原核对耗时－辅助核对与复核耗时”计算，同时扣除维护、标注、异常处置和培训投入；本轮不提供ROI或采购回本期。')
h('启动前必须补齐的事项')
p('业务负责人确认正式Brief、单批次范围及批准岗位；运营确认资源层级、互斥、缓冲与容量；质量确认样本标注、审计字段、访问范围、留存删除和异常升级；数字化确认三类工具的只读导出、编号映射与版本字段；项目负责人落实开发资源并澄清每周四小时的口径。')
h('建议的立项决定')
p('可以启动有边界的可行性验证，并以“只读数据关联、确定性冲突提示、版本及确认链”作为最小范围。完成真实样本复核和四周演练后，再决定轻量自建、二次开发或采购。当前证据不足以承诺全量上线时间、效率提升幅度或系统替换范围。')

page('参考资料 行业与产品')
p('以下链接均于2026年10月23日访问。正文使用方括号编号关联来源；原始网站及默认分支会继续变化。公开产品能力按厂商说明记录，适用性分析与实施建议由本报告推导。')
sources=[
('L1','本地业务需求发现报告','discovery-report.md','合成访谈与需求假设；引用f编号均为dlg-venue-schedule-f编号的简写。'),
('S1','中国贸促会 中国展览经济发展报告2025发布','https://exhibition.ccpit.org/articles/6397','发布页2026-01-28；统计覆盖2025年。'),
('S1a','中国贸促会 中国展览经济发展报告2025中英文下载页','https://exhibition.ccpit.org/articles/6391','2026-01-26；本报告量化引用依据S1发布说明。'),
('S2','上海MICE产业高质量发展三年行动方案','https://desk.sh.gov.cn/zxfw/zcfg/dfzh/20260603/3f9cf5e437ce4a5b94645a2d84a895a7.html','沪商会展〔2026〕143号；2026-05-26印发，页面2026-06-03。'),
('S3','上海市政府 第十二届上交会成绩单','https://www.shanghai.gov.cn/nw31406/20260615/a1bae05632b74d8d953331ee35ceb540.html','页面正文日期2026-06-13；数据为该次活动。'),
('S4','Momentus Answers 官方平台能力说明','https://gomomentus.com/momentus-answers','页面注明2026年5月更新；引用Enterprise能力。'),
('S5','Skedda Booking Requests and Approvals','https://support.skedda.com/en/articles/11774950-booking-requests-approvals','官方帮助文档2026-07-07；含套餐与请求状态限制。'),
('S6','FullCalendar License','https://fullcalendar.io/license','Standard与Premium分别授权；部署前按选定版本复核。'),
]
for no,title,url,note in sources:
    if no=='L1':
        p(f'[{no}] {title}');p(f'仓库来源：data/source/{url}。{note}')
    else:link(f'[{no}] {title}',url);p(note)

page('参考资料 GitHub与核验记录')
sources=[
('G1','LibreBooking 当前仓库与README','https://github.com/LibreBooking/librebooking','元数据：PHP、GPL-3.0、未归档；README及安全提示已读取。'),
('G2','pretalx README','https://github.com/pretalx/pretalx/blob/main/README.rst','核实版本发布、技术栈与插件；授权描述与LICENSE不一致。'),
('G3','OpenProject README','https://github.com/opf/openproject/blob/dev/README.md','核实项目管理、社区版与企业扩展；未验证各功能的具体版本边界。'),
('G4','Eventyay README','https://github.com/fossasia/eventyay/blob/dev/README.rst','核实统一平台、技术栈、分支与插件说明；元数据标识Apache-2.0。'),
('G5','FullCalendar README','https://github.com/fullcalendar/fullcalendar/blob/main/README.md','核实组件定位、集成与项目结构；高级功能许可证见S6。'),
('G6','LibreBooking 冲突识别源码','https://github.com/LibreBooking/librebooking/blob/develop/lib/Application/Reservation/ReservationConflictIdentifier.php','本轮读取1至170行；文件SHA 3bab4c2fb8e2db013c5f9d14a36522ca251ed59f。'),
('G7','pretalx 排期版本模型源码','https://github.com/pretalx/pretalx/blob/main/src/pretalx/schedule/models/schedule.py','本轮读取1至125行；文件SHA a2250c654c616678a59e11dbd0d8469103006530。'),
('G8','pretalx 当前LICENSE','https://github.com/pretalx/pretalx/blob/main/LICENSE','本轮读取前80行；文件SHA 0276f493966e7b9b831030a8c3ba1d35961db404。')
]
for no,title,url,note in sources:link(f'[{no}] {title}',url);p(note)
p('核验深度说明：已做仓库元数据和文档审阅，以及G6、G7的限定范围源码检查；没有执行安装、漏洞扫描、性能测试或业务验收。许可证信息用于选型记录，实际复用以锁定版本的完整授权文件和依赖清单为准。')

for element in [doc._element,doc.styles.element]:
    for borders in list(element.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
doc.core_properties.title='展会排期与需求变更项目研究报告'
doc.core_properties.subject='行业现状 GitHub项目对比与试点建议'
doc.core_properties.author='项目研究'
out=ROOT/'.research/work/report_base.docx'
doc.save(out)
print(out)
