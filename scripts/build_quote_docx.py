"""Create the printable customer quote from data/exhibition-quote.json.

This is a reference budget, not an issued contract. Excel edits are not imported
automatically: update the shared JSON to regenerate both canonical outputs.
"""
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
MODEL = json.loads((ROOT / 'data/exhibition-quote.json').read_text(encoding='utf-8'))
OUTPUT = ROOT / 'docs/reports/上海展会客户全流程报价单_参考版_20261023.docx'
FONT = 'Microsoft YaHei'

def amount(row):
    q = MODEL['area'] if row[4] == 'area' else row[4]
    n = MODEL['show_days'] if row[5] == 'show_days' else row[5]
    if row[6] is None:
        return None
    return (Decimal(str(q)) * Decimal(str(n)) * Decimal(str(row[6]))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def totals():
    result = {name: sum((amount(row) for row in MODEL['items'] if row[7] == name), Decimal(0)) for name in ['本方收款', '客户直付']}
    return result['本方收款'], result['客户直付']

def money(n):
    return f'{n:,.2f}'

def fonts(run, size=10, bold=False, color='243B50'):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + key), FONT)

def paragraph(doc, text='', size=10, bold=False, after=4, color='243B50'):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(snap)
    fonts(p.add_run(text), size, bold, color)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    fonts(p.add_run(text), 15 if level == 1 else 11, True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(5)
    return p

def page_heading(doc, text):
    p = heading(doc, text)
    p.paragraph_format.page_break_before = True
    return p

def table(doc, headers, rows, widths, size=9.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell, width in zip(t.columns, widths):
        cell.width = Cm(width)
    for idx, row in enumerate([headers, *rows]):
        cells = t.rows[0].cells if idx == 0 else t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for col, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = Cm(width)
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '234A68' if idx == 0 else ('F1F5F9' if idx % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:' + side); e.set(qn('w:w'), '45'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcpr.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(snap)
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fonts(p.add_run(str(value)), size, idx == 0, 'FFFFFF' if idx == 0 else '243B50')
    paragraph(doc, '', after=1, size=2)
    return t

def hyperlink(doc, label, url):
    p = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); rf = OxmlElement('w:rFonts')
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + key), FONT)
    rp.append(rf); color = OxmlElement('w:color'); color.set(qn('w:val'), '234A68'); rp.append(color)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '19'); rp.append(sz)
    r.append(rp); tx = OxmlElement('w:t'); tx.text = label; r.append(tx); link.append(r); p._p.append(link)
    p.paragraph_format.space_after = Pt(6)

def detailed(doc, rows):
    data = []
    for row in rows:
        q = MODEL['area'] if row[4] == 'area' else row[4]
        n = MODEL['show_days'] if row[5] == 'show_days' else row[5]
        data.append([row[0], row[2] + '\n' + row[9], f'{q} × {n}\n{row[3]}', '待询价' if row[6] is None else money(row[6]), '未选' if row[7] == '不选' else money(amount(row)), row[7], row[8]])
    table(doc, ['编号', '项目及交付边界', '数量与次数', '单价 元', '金额 元', '收款方式', '依据'], data, [.9, 6.5, 2.1, 1.9, 2.2, 2.4, 1.4], size=8.6)

def create():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65); sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(.65); sec.footer_distance = Cm(.65)
    for style in doc.styles:
        if style.type == 1 or style.type == 2:
            style.font.name = FONT
            rf = style.element.get_or_add_rPr().get_or_add_rFonts()
            for k in list(rf.attrib):
                del rf.attrib[k]
            for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
                rf.set(qn('w:' + key), FONT)
    doc.styles['Normal'].font.size = Pt(10)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
    doc.styles['Normal'].element.get_or_add_pPr().append(snap)
    title = doc.add_paragraph(style='Title')
    fonts(title.add_run(MODEL['title']), 23, True, '000000')
    paragraph(doc, f"报价编号 {MODEL['number']}    编制日期 {MODEL['date']}", size=9.5, color='62768A')
    paragraph(doc, '预算参考版  待双方确认后对外定稿', bold=True, color='A4491A')
    paragraph(doc, '本报价覆盖展会客户从需求确认、场地预订、设计搭建、进场、展中服务到撤展和结算的业务费用。上海为已确认城市；面积、展期、展品及具体展会尚未确定，当前采用可调整示例。')
    table(doc, ['客户及项目', '填写信息', '报价方及商务条件', '填写信息'], [
        ['客户单位', '________________', '报价单位', '________________'],
        ['展会及展位号', '待确认', '客户授权代表', '________________'],
        ['场馆及具体档期', '待确认', '报价负责人', '________________'],
        ['示例规模', f"{MODEL['area']}㎡ 单层模块化光地展台", '报价有效期', '待双方约定'],
        ['示例时长', f"展览{MODEL['show_days']}天 进场{MODEL['setup_days']}天 撤场{MODEL['teardown_days']}天", '币种及税费', '人民币 税费待逐项核实']
    ], [3, 5.7, 3.5, 5.2], size=9.3)
    own, direct = totals(); total = own + direct; dep = Decimal(MODEL['deposit'])
    heading(doc, '预算汇总')
    table(doc, ['费用类别', '参考金额 元', '支付口径'], [
        ['本方收款预算', money(own), '规划设计 搭建 展品物流及展中服务等'],
        ['客户直付预算', money(direct), '按主办 场馆及第三方订单直接支付'],
        ['已选全流程预算', money(total), '不含未选项目及可退押金 税费待核'],
        ['可退押金资金占用参考', money(dep), '假设客户直接缴纳场馆 不列作服务成本'],
        ['安排资金参考', money(total + dep), '已选费用加押金 非最终应付账单']
    ], [6.2, 3.5, 7.7], size=9.8)
    paragraph(doc, MODEL['tax_note'], size=9, color='62768A')
    heading(doc, '本方服务款建议付款节点')
    rates = [Decimal(str(x)) for x in MODEL['payment_rates']]
    payments = [(own * rates[i]).quantize(Decimal('.01'), rounding=ROUND_HALF_UP) for i in range(2)]
    payments.append(own - sum(payments))
    table(doc, ['节点', '建议比例', '参考金额 元', '触发条件'], [
        ['首期款', f'{rates[0]:.0%}', money(payments[0]), '双方确认报价及图纸后'],
        ['进场前款', f'{rates[1]:.0%}', money(payments[1]), '进场前双方约定日期'],
        ['验收尾款', f'{rates[2]:.0%}', money(payments[2]), '验收后按最终结算重算']
    ], [3.5, 2.3, 3.5, 8.1])
    paragraph(doc, '比例与付款时间均为建议，客户直付项目按其订单付款。最终服务尾款＝双方确认结算金额－本方服务款累计已付。', size=9)

    page_heading(doc, '展前策划与制作搭建费用')
    paragraph(doc, '单价及数量均可调整。金额按数量×次数或天数×单价计算；按展期或包干计价的项目不再乘展览天数。未选项目仅供增补选择。', size=9.4)
    detailed(doc, MODEL['items'][:14])
    paragraph(doc, 'Q06已包括搭建及拆除人工；Q14仅列展材运输。Q13只用于基础软膜之外的独立画面。费用范围变更时须保留原报价并出具新版本。', size=9.2)

    page_heading(doc, '场馆手续与现场执行费用')
    paragraph(doc, '场馆价格来自已公布展会的参考标准。目标展会及具体场馆未确定前，管理费、电箱、证件和审核要求须重新核实。', size=9.4)
    detailed(doc, MODEL['items'][14:28])
    paragraph(doc, 'Q20为客户展品运输，与Q14搭建展材运输区分。Q21装卸如已由主办方包付，应取消本项。Q28为普通废弃物清运，避免与Q06拆除人工重复。', size=9.2)

    page_heading(doc, '可选项目及商务约定')
    paragraph(doc, '以下项目默认未选，不进入第1页预算汇总。Q13额外画面和Q26摄影亦为未选项目，列在对应流程明细内。启用前须确认收款主体和价格。', size=9.4)
    detailed(doc, MODEL['items'][28:])
    heading(doc, '费用范围和变更约定')
    for text in [
        '当前方案适用于普通单层模块化展台、常规工作时段、轻型展品及上海同城运输。双层、超高、重型设备、危险品、远距离运输及特殊吊装另行核价。',
        '切换为标准展位套餐时，应核对主办套餐已含的围板、桌椅、灯具、地毯和电源，删除对应重复项目后再确认；不能只替换场租一项。',
        '图纸、画面、材料、设备功率、物料数量和进撤场时间以双方确认版本为准。修改导致费用或排期变化时，应列明增减项及证据，经双方授权人员批准后执行。',
        '现场加班费用可能同时涉及场馆和人员，Q29仅是特定场馆时段的参考，不含人员加班。实际加班需事先取得场馆许可并重新报价。',
        '取消、延期、缩减面积及已采购材料的处理条件由双方在正式合同中约定；本参考版不预设退款比例、违约金或场地保留期限。',
        '发票类型、含税单价、报价有效期、现场响应时段、最终验收日期、付款账户及押金退款条件，均需在对外定稿前补充确认。'
    ]:
        paragraph(doc, text, size=9.8)
    heading(doc, '对外发送前核对')
    paragraph(doc, '确认展会及场馆、实际面积及开口、展期及进撤场安排、客户展品参数、图纸清单、逐项含税报价、客户授权代表和有效期。未报价项目不得以零元替代。', size=10)

    page_heading(doc, '全流程交付与人工确认')
    table(doc, ['流程节点', '主要交付及验收依据', '人工确认角色'], [[r[0], r[1] + '\n' + r[3], r[2]] for r in MODEL['milestones']], [3.3, 9.5, 4.6], size=9.1)
    heading(doc, '需求变更记录模板')
    table(doc, ['记录字段', '待填写内容'], [
        ['变更编号及原报价版本', '________________________________________'],
        ['变更内容及原因依据', '________________________________________'],
        ['本方费用增减及排期影响', '人民币 __________ 元    排期影响 __________'],
        ['客户及本方授权确认', '客户 __________    本方 __________    日期 __________'],
        ['执行版本及证据位置', '________________________________________']
    ], [5.6, 11.8], size=9.2)
    paragraph(doc, 'Excel中只有“已批准”的本方费用差额计入修订预算；状态仅为登记，不替代签字与权限核验。客户直付变化直接更新明细并保存新版本，避免再计入本方变更差额。', size=9.2)
    heading(doc, '最终结算与押金跟踪')
    table(doc, ['结算项目', '金额 元', '确认依据'], [
        ['最终本方服务结算金额', '待填写', '双方签字结算 含已批准变更'],
        ['本方累计已付及应付或应退', '待填写', '付款凭证与结算余额'],
        ['场馆押金实缴 已退及扣款', '待填写', '场馆收据 退款凭据及扣款依据'],
        ['押金待退余额', '待填写', '实缴－已退－场馆确认扣款']
    ], [6.5, 2.7, 8.2], size=9.1)
    paragraph(doc, '押金按客户直接支付场馆的假设设计，服务尾款与场馆退押金分开核对；同一笔押金扣款不可再重复列入本方服务结算。', size=9.2)

    page_heading(doc, '参考价格依据与确认页')
    paragraph(doc, '以下资料于2026年10月23日核查。公开价格仅是特定展会、城市和服务规格下的锚点；本报价还含编制人的预算假设，未取得目标项目供应商正式报价。来源未统一注明税费，当前不做税率推定或换算。')
    for source in MODEL['sources']:
        heading(doc, source['id'] + ' ' + source['title'], 2)
        paragraph(doc, source['note'], size=9.6)
        if source['url']:
            hyperlink(doc, '打开官方资料', source['url'])
    heading(doc, '文件使用与信息保护')
    paragraph(doc, 'Excel可调整参数、选用方式、单价、变更状态和结算实收。Word是当前参数的静态报价快照，Excel修改后应同步修订Word及版本号，避免客户与执行人员使用不同版本。', size=9.5)
    paragraph(doc, '客户名单、手机号、证件、发票税号和银行资料按办理所需另行受控传递，限制访问人员并约定保留期限。公开仓库中仅保留空白或示例字段，不放入真实客户信息。', size=9.5)
    heading(doc, '双方确认')
    table(doc, ['确认事项', '签字或填写'], [
        ['客户单位及授权代表', '________________________________________'],
        ['报价单位及负责人', '________________________________________'],
        ['定稿报价版本及有效期', '________________________________________'],
        ['含税价格及开票条件', '________________________________________'],
        ['确认日期', '________________________________________']
    ], [5.7, 11.7], size=9.4)
    paragraph(doc, '本参考版未签字，不代表场地已锁档、施工图已批准或双方已签订正式合同。', size=9.2, color='62768A')

    hp = sec.header.paragraphs[0]; fonts(hp.add_run('DHLS  展会客户业务报价  上海参考版'), 8.5, color='62768A')
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fonts(fp.add_run(MODEL['number'] + '    第 '), 8, color='62768A')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); fp._p.append(field)
    fonts(fp.add_run(' 页'), 8, color='62768A')
    doc.core_properties.title = MODEL['title']; doc.core_properties.subject = '展会客户全流程参考预算及结算模板'
    doc.core_properties.author = 'DHLS'
    # Normalize unused table styles and document defaults as well as visible runs.
    for rf in doc.styles.element.iter(qn('w:rFonts')):
        for key in list(rf.attrib):
            del rf.attrib[key]
        for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            rf.set(qn('w:' + key), FONT)
    doc.save(OUTPUT)
    print(json.dumps({'output': str(OUTPUT), 'self': float(own), 'direct': float(direct), 'total': float(total)}, ensure_ascii=True))

if __name__ == '__main__':
    create()
